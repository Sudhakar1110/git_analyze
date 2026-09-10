import frappe
from frappe import _
import time
import threading
import traceback
import json
import re
import base64
from frappe.utils import now_datetime, cint

try:
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

DECOMMISSIONED_MODELS = [
    "llama3-8b-8192", "llama-3.1-8b-instant", "llama3-70b-8192",
    "llama-3.3-70b-versatile", "mixtral-8x7b-32768", "gemma2-9b-it",
]

ANALYSIS_LOCK = threading.Lock()

SECTION_LABELS = {
    "purpose": "Executive Summary",
    "tech_stack": "Technology Stack",
    "architecture": "System Architecture",
    "entry_points": "Entry Points & Boot Sequence",
    "key_modules": "Core Modules & Components",
    "data_flow": "Data Flow & Pipeline",
    "api_endpoints": "API Endpoints & Contracts",
    "database_models": "Data Models & Schema",
    "dependencies": "Dependency Matrix",
    "how_to_run": "Deployment & Setup Guide",
}

SECTION_ICONS = {
    "purpose": "01",
    "tech_stack": "02",
    "architecture": "03",
    "entry_points": "04",
    "key_modules": "05",
    "data_flow": "06",
    "api_endpoints": "07",
    "database_models": "08",
    "dependencies": "09",
    "how_to_run": "10",
}

SECTION_ORDER = [
    "purpose", "tech_stack", "architecture", "entry_points", "key_modules",
    "data_flow", "api_endpoints", "database_models", "dependencies", "how_to_run",
]


def has_permission(doc, user):
    if user == "Administrator":
        return True
    return frappe.has_permission(doc.doctype, "read", user=user)


def run_analysis_on_submit(doc, method):
    pass


@frappe.whitelist()
def analyze_repo(github_url, branch="main", depth="medium", questions=""):
    if not github_url:
        frappe.throw(_("GitHub URL is required"))
    if "github.com" not in github_url:
        frappe.throw(_("Please enter a valid GitHub URL"))

    settings = frappe.get_doc("Analysis Settings", "Analysis Settings")
    if not settings.groq_api_key:
        frappe.throw(_("Please configure Groq API key in Analysis Settings"))

    max_files = {"shallow": 50, "medium": 100, "deep": 200}.get(depth, 100)

    from git_analyze.github_fetcher import GitHubFetcher
    fetcher_temp = GitHubFetcher()
    owner, repo_name = fetcher_temp.parse_github_url(github_url)

    repo_analysis = frappe.get_doc({
        "doctype": "Repo Analysis",
        "github_url": github_url,
        "repo_name": f"{owner}/{repo_name}",
        "branch": branch,
        "status": "In Progress",
    })
    repo_analysis.insert(ignore_permissions=True)
    frappe.db.commit()

    site_name = frappe.conf.site_name or frappe.local.site

    t = threading.Thread(
        target=_bg_run,
        args=(site_name, repo_analysis.name, github_url, branch, max_files),
        daemon=True,
        name=f"analysis-{repo_analysis.name}",
    )
    t.start()

    return {"status": "success", "repo_analysis": repo_analysis.name, "name": repo_analysis.name}


def _bg_run(site, name, url, branch, max_files):
    try:
        frappe.init(site=site, force=True)
        frappe.connect()
        frappe.set_user("Administrator")
        _do_analysis(name, url, branch, max_files)
    except Exception as e:
        error_msg = f"{type(e).__name__}: {str(e)}"
        try:
            frappe.db.set_value("Repo Analysis", name, "status", "Failed")
            frappe.db.set_value("Repo Analysis", name, "full_output", error_msg)
            frappe.db.commit()
        except Exception:
            pass
        try:
            frappe.log_error(f"BG Analysis failed for {name}: {error_msg}\n{traceback.format_exc()}")
        except Exception:
            pass
    finally:
        try:
            frappe.db.close()
        except Exception:
            pass
        try:
            frappe.destroy()
        except Exception:
            pass


def _do_analysis(name, github_url, branch, max_files):
    start_time = time.time()
    settings = frappe.get_doc("Analysis Settings", "Analysis Settings")

    frappe.db.set_value("Repo Analysis", name, "full_output", "Fetching repository files...")
    frappe.db.commit()

    from git_analyze.github_fetcher import GitHubFetcher
    fetcher = GitHubFetcher(github_token=settings.get_github_token())
    repo_data = fetcher.fetch_repository(
        github_url=github_url,
        branch=branch,
        max_files=max_files,
        skip_patterns=settings.get_skip_patterns(),
    )

    if not repo_data.get("files"):
        frappe.db.set_value("Repo Analysis", name, "status", "Failed")
        frappe.db.set_value("Repo Analysis", name, "full_output", "No files found. Check URL and branch.")
        frappe.db.commit()
        return

    groq_model = settings.groq_model or "openai/gpt-oss-20b"
    if groq_model in DECOMMISSIONED_MODELS:
        groq_model = "openai/gpt-oss-20b"
        frappe.db.set_value("Analysis Settings", "Analysis Settings", "groq_model", groq_model)
        frappe.db.commit()

    frappe.db.set_value("Repo Analysis", name, "full_output", f"Analyzing {repo_data['analyzed_files']} files with AI...")
    frappe.db.commit()

    from git_analyze.groq_client import GroqClient
    groq_client = GroqClient(
        api_key=settings.get_groq_api_key(),
        model=groq_model,
    )

    result = groq_client.analyze_repository(
        repo_name=f"{repo_data['owner']}/{repo_data['repo']}",
        branch=repo_data["branch"],
        file_structure=repo_data["structure"],
        file_contents=repo_data["files"],
        language=settings.output_language,
    )

    sections = groq_client.parse_analysis_sections(result["content"])
    analysis_time = time.time() - start_time

    frappe.db.set_value("Repo Analysis", name, "status", "Completed")
    frappe.db.set_value("Repo Analysis", name, "analysis_date", now_datetime())
    frappe.db.set_value("Repo Analysis", name, "purpose", sections.get("purpose", ""))
    frappe.db.set_value("Repo Analysis", name, "tech_stack", sections.get("tech_stack", ""))
    frappe.db.set_value("Repo Analysis", name, "architecture", sections.get("architecture", ""))
    frappe.db.set_value("Repo Analysis", name, "entry_points", sections.get("entry_points", ""))
    frappe.db.set_value("Repo Analysis", name, "key_modules", sections.get("key_modules", ""))
    frappe.db.set_value("Repo Analysis", name, "data_flow", sections.get("data_flow", ""))
    frappe.db.set_value("Repo Analysis", name, "api_endpoints", sections.get("api_endpoints", ""))
    frappe.db.set_value("Repo Analysis", name, "database_models", sections.get("database_models", ""))
    frappe.db.set_value("Repo Analysis", name, "dependencies", sections.get("dependencies", ""))
    frappe.db.set_value("Repo Analysis", name, "how_to_run", sections.get("how_to_run", ""))
    frappe.db.set_value("Repo Analysis", name, "full_output", result["content"])
    frappe.db.set_value("Repo Analysis", name, "groq_model", result["model"])
    frappe.db.set_value("Repo Analysis", name, "token_usage", result["token_usage"]["total_tokens"])
    frappe.db.set_value("Repo Analysis", name, "analysis_time", analysis_time)
    frappe.db.set_value("Repo Analysis", name, "file_count", repo_data["analyzed_files"])
    frappe.db.commit()


@frappe.whitelist()
def save_settings(groq_api_key=None, groq_model=None, max_files_limit=None, output_language=None, github_token=None):
    settings = frappe.get_doc("Analysis Settings", "Analysis Settings")
    if groq_api_key is not None:
        settings.groq_api_key = groq_api_key
    if groq_model is not None:
        settings.groq_model = groq_model
    if max_files_limit is not None:
        settings.max_files_limit = int(max_files_limit)
    if output_language is not None:
        settings.output_language = output_language
    if github_token is not None:
        settings.github_token = github_token
    settings.save(ignore_permissions=True)
    frappe.db.commit()
    return {"status": "success"}


@frappe.whitelist()
def ask_followup(repo_analysis_name, question):
    if not repo_analysis_name or not question:
        frappe.throw(_("Both repo_analysis_name and question are required"))
    settings = frappe.get_doc("Analysis Settings", "Analysis Settings")
    repo_analysis = frappe.get_doc("Repo Analysis", repo_analysis_name)
    from git_analyze.groq_client import GroqClient
    groq_client = GroqClient(api_key=settings.get_groq_api_key(), model=settings.groq_model)
    result = groq_client.ask_followup(
        repo_name=repo_analysis.repo_name,
        question=question,
        previous_analysis=repo_analysis.full_output or "",
        language=settings.output_language,
    )
    return {"status": "success", "answer": result["content"]}


@frappe.whitelist()
def get_analysis_status(repo_analysis_name):
    repo_analysis = frappe.get_doc("Repo Analysis", repo_analysis_name)
    return {
        "status": repo_analysis.status,
        "repo_name": repo_analysis.repo_name,
        "token_usage": repo_analysis.token_usage,
        "analysis_time": repo_analysis.analysis_time,
        "file_count": repo_analysis.file_count,
        "full_output": repo_analysis.full_output,
    }


@frappe.whitelist()
def get_analysis_results(repo_analysis_name):
    r = frappe.get_doc("Repo Analysis", repo_analysis_name)
    return {
        "name": r.name, "github_url": r.github_url, "repo_name": r.repo_name,
        "branch": r.branch, "status": r.status, "purpose": r.purpose,
        "tech_stack": r.tech_stack, "architecture": r.architecture,
        "entry_points": r.entry_points, "key_modules": r.key_modules,
        "data_flow": r.data_flow, "api_endpoints": r.api_endpoints,
        "database_models": r.database_models, "dependencies": r.dependencies,
        "how_to_run": r.how_to_run, "full_output": r.full_output,
        "groq_model": r.groq_model, "token_usage": r.token_usage,
        "analysis_time": r.analysis_time, "file_count": r.file_count,
    }


# ──────────────────────────────────────────────────────────────────────────────
# REPORT DATA HELPERS
# ──────────────────────────────────────────────────────────────────────────────

def _get_analysis_data(name):
    r = frappe.get_doc("Repo Analysis", name)
    if r.status != "Completed" or not r.full_output:
        frappe.throw(_("Analysis is not completed yet. Export is available only for completed analyses."))
    return r


def _collect_sections(r):
    data = {}
    for key in SECTION_ORDER:
        val = getattr(r, key, None)
        if val and val.strip():
            data[key] = val.strip()
    return data


def _esc(text):
    s = str(text or "")
    s = s.replace("&", "&amp;")
    s = s.replace("<", "&lt;")
    s = s.replace(">", "&gt;")
    s = s.replace('"', "&quot;")
    s = s.replace("'", "&#x27;")
    return s


def _md_to_html(text):
    if not text:
        return ""
    lines = text.split("\n")
    out = []
    in_code = False
    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            out.append("<br>" if in_code else "")
            continue
        if in_code:
            out.append(f'<code>{_esc(line)}</code><br>')
            continue
        escaped = _esc(line)
        if escaped.startswith("### "):
            out.append(f'<h4>{escaped[4:]}</h4>')
        elif escaped.startswith("## "):
            out.append(f'<h3>{escaped[3:]}</h3>')
        elif escaped.startswith("# "):
            out.append(f'<h2>{escaped[2:]}</h2>')
        elif escaped.startswith("- ") or escaped.startswith("* "):
            out.append(f'<div style="padding-left:16px;">&#9654; {escaped[2:]}</div>')
        elif escaped.startswith("| ") or escaped.startswith("|--"):
            out.append(f'<div style="font-family:monospace;font-size:12px;">{escaped}</div>')
        elif escaped.strip() == "":
            out.append("<br>")
        else:
            bolded = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', escaped)
            out.append(f'<div>{bolded}</div>')
    return "\n".join(out)


# ──────────────────────────────────────────────────────────────────────────────
# PROFESSIONAL HTML / PDF REPORT
# ──────────────────────────────────────────────────────────────────────────────

def _build_professional_html(r):
    sections = _collect_sections(r)
    now = str(r.analysis_date or r.creation or now_datetime())
    score = _complexity_score(r)

    nav_items = ""
    content_sections = ""
    toc_items = ""
    for i, key in enumerate(SECTION_ORDER):
        if key not in sections:
            continue
        label = SECTION_LABELS.get(key, key)
        num = SECTION_ICONS.get(key, f"{i+1:02d}")
        anchor = f"section-{key}"
        toc_items += f'<li><a href="#{anchor}" style="color:#475569;text-decoration:none;font-size:13px;padding:6px 0;display:block;border-bottom:1px dashed #e2e8f0;">{num}. {label}</a></li>'
        content_sections += f"""
        <div id="{anchor}" style="page-break-inside:avoid;margin-bottom:40px;">
            <div style="display:flex;align-items:center;gap:12px;margin-bottom:16px;">
                <div style="width:36px;height:36px;background:#10b981;color:#fff;border-radius:8px;display:flex;align-items:center;justify-content:center;font-size:13px;font-weight:700;flex-shrink:0;">{num}</div>
                <h2 style="color:#0f172a;font-size:20px;font-weight:700;margin:0;">{label}</h2>
            </div>
            <div style="background:#fff;border:1px solid #e2e8f0;border-radius:12px;padding:24px;margin-left:48px;box-shadow:0 1px 3px rgba(0,0,0,0.04);">
                <div style="color:#334155;font-size:13px;line-height:1.8;">{_md_to_html(sections[key])}</div>
            </div>
        </div>"""

    tech_badges = ""
    if r.tech_stack:
        techs = re.findall(r'`([^`]+)`', r.tech_stack)
        if not techs:
            techs = re.findall(r'\b(Python|JavaScript|TypeScript|React|Vue|Angular|Node|Django|Flask|FastAPI|Express|Next|Nuxt|Laravel|Rails|Spring|Go|Rust|Java|C\+\+|PHP|Ruby|HTML|CSS|SQL|PostgreSQL|MySQL|MongoDB|Redis|Docker|Kubernetes|AWS|GCP|Azure|Tailwind|Bootstrap|Webpack|Vite)\b', r.tech_stack, re.I)
        for t in techs[:12]:
            tech_badges += f'<span style="display:inline-block;padding:4px 12px;background:#ecfdf5;color:#059669;border-radius:20px;font-size:11px;font-weight:600;margin:2px;">{t}</span>'

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>Repository Analysis Report — {r.repo_name}</title>
<style>
@page {{ size: A4; margin: 0; }}
* {{ margin: 0; padding: 0; box-sizing: border-box; }}
body {{ font-family: 'Segoe UI', 'Helvetica Neue', Arial, sans-serif; color: #1e293b; background: #f8fafc; line-height: 1.6; }}
.cover {{ background: linear-gradient(135deg, #0f172a 0%, #1e293b 40%, #0f766e 100%); color: #fff; min-height: 100vh; display: flex; flex-direction: column; justify-content: center; padding: 60px 80px; page-break-after: always; }}
.cover-brand {{ display: flex; align-items: center; gap: 12px; margin-bottom: 48px; }}
.cover-brand-icon {{ width: 48px; height: 48px; background: #10b981; border-radius: 12px; display: flex; align-items: center; justify-content: center; font-size: 24px; }}
.cover-brand-text {{ font-size: 14px; font-weight: 600; letter-spacing: 0.1em; text-transform: uppercase; opacity: 0.8; }}
.cover h1 {{ font-size: 42px; font-weight: 800; line-height: 1.15; margin-bottom: 16px; }}
.cover h1 span {{ color: #34d399; }}
.cover-repo {{ font-size: 20px; opacity: 0.9; margin-bottom: 40px; font-weight: 300; }}
.cover-meta {{ display: grid; grid-template-columns: 1fr 1fr 1fr 1fr; gap: 20px; margin-bottom: 48px; }}
.cover-meta-item {{ background: rgba(255,255,255,0.08); border-radius: 12px; padding: 20px; border: 1px solid rgba(255,255,255,0.1); }}
.cover-meta-label {{ font-size: 11px; text-transform: uppercase; letter-spacing: 0.08em; opacity: 0.6; margin-bottom: 4px; }}
.cover-meta-value {{ font-size: 22px; font-weight: 700; }}
.cover-footer {{ border-top: 1px solid rgba(255,255,255,0.15); padding-top: 24px; display: flex; justify-content: space-between; font-size: 13px; opacity: 0.7; }}
.page {{ padding: 48px 60px; max-width: 900px; margin: 0 auto; background: #fff; }}
.page-toc {{ page-break-after: always; }}
.page-toc h2 {{ font-size: 24px; font-weight: 700; color: #0f172a; margin-bottom: 24px; padding-bottom: 12px; border-bottom: 3px solid #10b981; }}
.page-toc ol {{ padding-left: 24px; }}
.page-toc li {{ list-style: none; padding: 0; }}
.section-header {{ display: flex; align-items: center; gap: 14px; margin-bottom: 20px; page-break-after: avoid; }}
.section-num {{ width: 40px; height: 40px; background: linear-gradient(135deg, #10b981, #059669); color: #fff; border-radius: 10px; display: flex; align-items: center; justify-content: center; font-size: 15px; font-weight: 700; flex-shrink: 0; }}
.section-title {{ font-size: 22px; font-weight: 700; color: #0f172a; }}
.section-body {{ background: #fff; border: 1px solid #e2e8f0; border-radius: 12px; padding: 28px; margin-bottom: 36px; box-shadow: 0 1px 3px rgba(0,0,0,0.04); margin-left: 54px; page-break-inside: avoid; }}
.section-content {{ color: #334155; font-size: 13.5px; line-height: 1.85; }}
.tech-badges {{ display: flex; flex-wrap: wrap; gap: 6px; margin-top: 12px; }}
.score-bar {{ display: flex; align-items: center; gap: 12px; margin: 20px 0; }}
.score-track {{ flex: 1; height: 10px; background: #e2e8f0; border-radius: 5px; overflow: hidden; }}
.score-fill {{ height: 100%; border-radius: 5px; transition: width 0.3s; }}
.footer-report {{ text-align: center; padding: 32px; color: #94a3b8; font-size: 11px; border-top: 1px solid #e2e8f0; margin-top: 40px; }}
</style>
</head>
<body>

<!-- COVER PAGE -->
<div class="cover">
    <div class="cover-brand">
        <div class="cover-brand-icon">&#9881;</div>
        <div class="cover-brand-text">Git Analyzer &mdash; Automated Code Analysis Report</div>
    </div>
    <h1>Repository<br><span>Analysis Report</span></h1>
    <div class="cover-repo">{r.repo_name} &mdash; {r.branch} branch</div>
    <div class="cover-meta">
        <div class="cover-meta-item">
            <div class="cover-meta-label">Files Analyzed</div>
            <div class="cover-meta-value">{r.file_count or 0}</div>
        </div>
        <div class="cover-meta-item">
            <div class="cover-meta-label">AI Model</div>
            <div class="cover-meta-value" style="font-size:15px;">{r.groq_model or 'default'}</div>
        </div>
        <div class="cover-meta-item">
            <div class="cover-meta-label">Tokens Consumed</div>
            <div class="cover-meta-value">{r.token_usage or 0:,}</div>
        </div>
        <div class="cover-meta-item">
            <div class="cover-meta-label">Complexity Score</div>
            <div class="cover-meta-value" style="color:#34d399;">{score}/100</div>
        </div>
    </div>
    <div class="cover-footer">
        <span>Report generated on {now}</span>
        <span>{r.github_url}</span>
    </div>
</div>

<!-- TABLE OF CONTENTS -->
<div class="page page-toc">
    <h2>Table of Contents</h2>
    <ol>{toc_items}</ol>
    <div style="margin-top:32px;padding:20px;background:#f0fdf4;border-radius:8px;border-left:4px solid #10b981;">
        <strong style="color:#065f46;">About this report</strong>
        <p style="color:#047857;font-size:13px;margin-top:4px;">This document was automatically generated by Git Analyzer using AI-powered code analysis. It provides a comprehensive technical overview of the repository including architecture, dependencies, workflows, and deployment guidance.</p>
    </div>
</div>

<!-- CONTENT SECTIONS -->
<div class="page">
    {content_sections}

    <div class="footer-report">
        <strong>Git Analyzer</strong> &mdash; Automated Repository Analysis<br>
        Report generated on {now} &nbsp;|&nbsp; AI Model: {r.groq_model or 'default'} &nbsp;|&nbsp; Analysis ID: {r.name}
    </div>
</div>

</body>
</html>"""


def _complexity_score(r):
    score = 30
    if r.file_count and r.file_count > 10:
        score += 10
    if r.file_count and r.file_count > 50:
        score += 10
    if r.file_count and r.file_count > 100:
        score += 10
    if r.token_usage and r.token_usage > 2000:
        score += 10
    if r.token_usage and r.token_usage > 5000:
        score += 10
    if r.tech_stack and len(r.tech_stack) > 200:
        score += 10
    if r.architecture and len(r.architecture) > 300:
        score += 10
    return min(score, 100)


# ──────────────────────────────────────────────────────────────────────────────
# MARKDOWN EXPORT
# ──────────────────────────────────────────────────────────────────────────────

def _build_professional_markdown(r):
    sections = _collect_sections(r)
    score = _complexity_score(r)
    now = str(r.analysis_date or r.creation or now_datetime())

    md = f"""# {r.repo_name} — Repository Analysis Report

> **Generated by** Git Analyzer &nbsp;|&nbsp; **Date:** {now} &nbsp;|&nbsp; **Report ID:** {r.name}

---

## Overview

| Metric | Value |
|--------|-------|
| **Repository** | {r.repo_name} |
| **Branch** | {r.branch} |
| **Source** | {r.github_url} |
| **Files Analyzed** | {r.file_count or 0} |
| **AI Model** | {r.groq_model or 'default'} |
| **Tokens Consumed** | {r.token_usage or 0:,} |
| **Analysis Time** | {r.analysis_time or 0}s |
| **Complexity Score** | {score}/100 |

---

## Table of Contents

"""
    for i, key in enumerate(SECTION_ORDER):
        if key in sections:
            md += f"{i+1}. [{SECTION_LABELS.get(key, key)}](#section-{key})\n"
    md += "\n---\n\n"

    for i, key in enumerate(SECTION_ORDER):
        if key not in sections:
            continue
        label = SECTION_LABELS.get(key, key)
        md += f"## {i+1}. {label}\n\n"
        md += f"{sections[key]}\n\n---\n\n"

    md += f"""## Appendix: Full AI Analysis Output

<details>
<summary>Click to expand full raw output</summary>

```
{r.full_output}
```

</details>

---

*Report generated by Git Analyzer on {now}*
"""
    return md


# ──────────────────────────────────────────────────────────────────────────────
# JSON EXPORT
# ──────────────────────────────────────────────────────────────────────────────

def _build_professional_json(r):
    sections = _collect_sections(r)
    return {
        "report": {
            "id": r.name,
            "generated_by": "Git Analyzer",
            "generated_at": str(now_datetime()),
            "version": "2.0",
        },
        "repository": {
            "name": r.repo_name,
            "url": r.github_url,
            "branch": r.branch,
            "status": r.status,
        },
        "analysis": {
            "files_analyzed": r.file_count or 0,
            "ai_model": r.groq_model or "default",
            "tokens_consumed": r.token_usage or 0,
            "time_seconds": r.analysis_time or 0,
            "complexity_score": _complexity_score(r),
            "date": str(r.analysis_date or r.creation),
        },
        "sections": {key: sections.get(key, "") for key in SECTION_ORDER},
        "section_labels": {key: SECTION_LABELS.get(key, key) for key in SECTION_ORDER},
        "full_output": r.full_output or "",
    }


# ──────────────────────────────────────────────────────────────────────────────
# WHITELISTED EXPORT ENDPOINTS
# ──────────────────────────────────────────────────────────────────────────────

def _save_as_file(file_data, filename, folder="Home"):
    file_doc = frappe.get_doc({
        "doctype": "File",
        "file_name": filename,
        "is_private": 1,
        "content": file_data,
        "folder": folder,
    })
    file_doc.save(ignore_permissions=True)
    frappe.db.commit()
    return file_doc.file_url


def _safe_get_analysis(name):
    try:
        r = frappe.get_doc("Repo Analysis", name)
        if r.status != "Completed" or not r.full_output:
            return None, "Analysis is not completed yet. Export is available only for completed analyses."
        return r, None
    except Exception as e:
        return None, f"Could not load analysis: {str(e)}"


@frappe.whitelist()
def export_as_markdown(repo_analysis_name):
    r, err = _safe_get_analysis(repo_analysis_name)
    if err:
        return {"error": err}
    return {"content": _build_professional_markdown(r), "filename": f"{r.repo_name.replace('/', '_')}_analysis.md"}


@frappe.whitelist()
def export_as_html(repo_analysis_name):
    r, err = _safe_get_analysis(repo_analysis_name)
    if err:
        return {"error": err}
    return {"content": _build_professional_html(r), "filename": f"{r.repo_name.replace('/', '_')}_analysis.html"}


@frappe.whitelist()
def export_as_json(repo_analysis_name):
    r, err = _safe_get_analysis(repo_analysis_name)
    if err:
        return {"error": err}
    data = _build_professional_json(r)
    return {"content": json.dumps(data, indent=2), "filename": f"{r.repo_name.replace('/', '_')}_analysis.json"}


@frappe.whitelist()
def export_as_pdf(repo_analysis_name):
    r, err = _safe_get_analysis(repo_analysis_name)
    if err:
        return {"error": err}
    try:
        html = _build_professional_html(r)
        from frappe.utils.pdf import get_pdf
        pdf = get_pdf(html, options={
            "page-size": "A4",
            "margin-top": "0mm",
            "margin-bottom": "0mm",
            "margin-left": "0mm",
            "margin-right": "0mm",
            "print-media-type": "",
            "enable-local-file-access": "",
        })
        filename = f"{r.repo_name.replace('/', '_')}_analysis.pdf"
        file_url = _save_as_file(pdf, filename, "Home")
        return {"file_url": file_url, "filename": filename}
    except ImportError:
        return {"error": "wkhtmltopdf is not installed on the server."}
    except Exception as e:
        return {"error": f"PDF generation failed: {str(e)}"}


@frappe.whitelist()
def export_as_docx(repo_analysis_name):
    r, err = _safe_get_analysis(repo_analysis_name)
    if err:
        return {"error": err}
    if not HAS_DOCX:
        return {"error": "python-docx is not installed. Run: bench pip install python-docx"}
    try:
        from docx import Document
        import io

        doc = Document()

        section = doc.sections[0]
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)

        for level in range(1, 4):
            hs = doc.styles[f"Heading {level}"]
            hs.font.color.rgb = RGBColor(15, 23, 42)
            hs.font.bold = True

        _add_docx_cover(doc, r)
        doc.add_page_break()

        _add_docx_toc_page(doc, r)
        doc.add_page_break()

        sections = _collect_sections(r)
        for i, key in enumerate(SECTION_ORDER):
            if key not in sections:
                continue
            label = SECTION_LABELS.get(key, key)
            num = SECTION_ICONS.get(key, f"{i+1:02d}")

            heading = doc.add_heading(f"{num}. {label}", level=1)
            for run in heading.runs:
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(15, 23, 42)

            body = doc.add_paragraph()
            body.paragraph_format.space_before = Pt(8)
            _add_styled_content(doc, sections[key])

            if i < len(SECTION_ORDER) - 1:
                doc.add_page_break()

        doc.add_page_break()
        _add_docx_appendix(doc, r)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        filename = f"{r.repo_name.replace('/', '_')}_analysis.docx"
        file_url = _save_as_file(buf.read(), filename, "Home")
        return {"file_url": file_url, "filename": filename}
    except ImportError:
        return {"error": "python-docx is not installed. Run: bench pip install python-docx"}
    except Exception as e:
        return {"error": f"DOCX generation failed: {str(e)}"}


def _add_docx_cover(doc, r):
    for _ in range(6):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("GIT ANALYZER")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(16, 185, 129)
    run.font.bold = True

    doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Repository Analysis Report")
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(r.repo_name)
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Branch: {r.branch}")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(100, 116, 139)

    for _ in range(4):
        doc.add_paragraph("")

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Report Date", str(r.analysis_date or r.creation)),
        ("AI Model", r.groq_model or "default"),
        ("Files Analyzed", str(r.file_count or 0)),
        ("Complexity Score", f"{_complexity_score(r)}/100"),
    ]
    for i, (k, v) in enumerate(meta):
        c0 = table.rows[i].cells[0]
        c1 = table.rows[i].cells[1]
        c0.text = k
        c1.text = v
        for cell in [c0, c1]:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(10)
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(100, 116, 139)
        c1.paragraphs[0].runs[0].font.bold = True


def _add_docx_toc_page(doc, r):
    h = doc.add_heading("Table of Contents", level=1)
    for run in h.runs:
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("")

    sections = _collect_sections(r)
    for i, key in enumerate(SECTION_ORDER):
        if key not in sections:
            continue
        label = SECTION_LABELS.get(key, key)
        num = SECTION_ICONS.get(key, f"{i+1:02d}")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(f"{num}.  {label}")
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("")
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = info.add_run("This report was automatically generated by Git Analyzer. It provides a comprehensive technical analysis of the repository covering architecture, dependencies, workflows, and deployment guidance.")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 116, 139)
    run.font.italic = True


def _add_docx_appendix(doc, r):
    h = doc.add_heading("Appendix: Full AI Analysis Output", level=1)
    for run in h.runs:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph("")

    output = r.full_output or "No output available."
    for para_text in output.split("\n"):
        if para_text.strip():
            p = doc.add_paragraph(para_text.strip())
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Consolas"
                run.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Generated by Git Analyzer — {now_datetime()}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(148, 163, 184)


def _add_styled_content(doc, text):
    if not text:
        return
    lines = text.split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(line[4:])
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        else:
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            p = doc.add_paragraph()
            run = p.add_run(clean)
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(51, 65, 85)


# ──────────────────────────────────────────────────────────────────────────────
# TEST CONNECTION
# ──────────────────────────────────────────────────────────────────────────────

@frappe.whitelist()
def test_groq_connection():
    import requests as req
    try:
        settings = frappe.get_doc("Analysis Settings", "Analysis Settings")
    except Exception:
        return {"status": "error", "error": "Analysis Settings not found"}
    if not settings.groq_api_key:
        return {"status": "error", "error": "No Groq API key configured"}

    try:
        api_key = settings.get_groq_api_key()
        model = settings.groq_model or "openai/gpt-oss-20b"
        if model in DECOMMISSIONED_MODELS:
            model = "openai/gpt-oss-20b"
            frappe.db.set_value("Analysis Settings", "Analysis Settings", "groq_model", model)
            frappe.db.commit()
        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
        payload = {"model": model, "messages": [{"role": "user", "content": "Say OK"}], "max_tokens": 5}
        r = req.post(url, json=payload, headers=headers, timeout=30)
        if r.status_code != 200:
            error_detail = r.json().get("error", {}).get("message", r.text[:200])
            return {"status": "error", "error": f"Groq API error ({r.status_code}): {error_detail}"}
        data = r.json()
        return {"status": "ok", "model": model, "response": data["choices"][0]["message"]["content"]}
    except req.exceptions.ConnectionError:
        return {"status": "error", "error": "Cannot reach api.groq.com"}
    except req.exceptions.Timeout:
        return {"status": "error", "error": "Connection timed out"}
    except Exception as e:
        return {"status": "error", "error": f"{type(e).__name__}: {str(e)}"}
